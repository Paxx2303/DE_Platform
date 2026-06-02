import pandas as pd
from openpyxl import load_workbook
import shutil
import os

txt_file = r"C:\Using\DE_Platform\sprint_1\1_business_analysis\deliverables\50_question.txt"
excel_v1 = r"C:\Using\DE_Platform\sprint_1\1_business_analysis\deliverables\AC_v1.xlsx"
excel_v2 = r"C:\Using\DE_Platform\sprint_1\1_business_analysis\deliverables\AC_v2.xlsx"

shutil.copyfile(excel_v1, excel_v2)

with open(txt_file, 'r', encoding='utf-8') as f:
    lines = f.readlines()

data = []
for line in lines:
    line = line.strip()
    if not line or line.startswith('|-'):
        continue
    if line.startswith('|'):
        line = line[1:]
    if line.endswith('|'):
        line = line[:-1]
    
    row = [cell.strip() for cell in line.split('|')]
    data.append(row)

wb = load_workbook(excel_v2)
if 'Sheet2' in wb.sheetnames:
    del wb['Sheet2']
ws = wb.create_sheet('Sheet2')

for row_idx, row_data in enumerate(data, start=1):
    for col_idx, value in enumerate(row_data, start=1):
        ws.cell(row=row_idx, column=col_idx, value=value)

wb.save(excel_v2)
print("Created AC_v2.xlsx successfully")
